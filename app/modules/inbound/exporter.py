"""按模板把入库单导出成 Excel（明细行展开，一张单占明细行数行）。

列顺序就是导出顺序；以后对方系统改模板，只改 COLUMNS 这一个列表。
"""
from __future__ import annotations

import sqlite3
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# (表头, 取值函数) —— row 是单据字段+明细字段的平铺字典
COLUMNS: list[tuple[str, object]] = [
    ("入库单号", lambda r: r["receipt_no"]),
    ("入库日期", lambda r: r["receipt_date"]),
    ("供应商", lambda r: r["supplier"]),
    ("物料类别", lambda r: r["category"]),
    ("品名", lambda r: r["name"]),
    ("规格", lambda r: r["spec"]),
    ("属性", lambda r: r["attr"]),
    ("数量", lambda r: r["qty"]),
    ("单位", lambda r: r["unit"]),
    ("单价", lambda r: r["unit_price"]),
    ("金额", lambda r: round((r["qty"] or 0) * (r["unit_price"] or 0), 2)),
    ("录入人", lambda r: r["clerk"]),
    ("定价人", lambda r: r["priced_by"]),
    ("定价时间", lambda r: r["priced_at"]),
    ("状态", lambda r: {"pending_pricing": "待定价", "priced": "已定价", "void": "作废"}.get(r["status"], r["status"])),
    ("明细备注", lambda r: r["note"]),
    ("单据备注", lambda r: r["receipt_note"]),
    ("来源", lambda r: "拍照识别" if r["source"] == "ocr" else "手工录入"),
]

_HEADER_FONT = Font(bold=True, color="FFFFFF")
_HEADER_FILL = PatternFill("solid", fgColor="305496")
_CENTER = Alignment(horizontal="center", vertical="center")


def query_rows(conn: sqlite3.Connection, start: str, end: str, statuses: list[str]) -> list[dict]:
    marks = ",".join("?" * len(statuses))
    sql = f"""
        SELECT i.*, r.receipt_no, r.receipt_date, r.supplier, r.source, r.status,
               r.clerk, r.note AS receipt_note
        FROM items i JOIN receipts r ON r.id = i.receipt_id
        WHERE r.status IN ({marks})
          AND ( ? = '' OR r.receipt_date >= ? )
          AND ( ? = '' OR r.receipt_date <= ? )
        ORDER BY r.receipt_date, r.id, i.seq
    """
    return [dict(x) for x in conn.execute(sql, (*statuses, start, start, end, end))]


def export_xlsx(conn: sqlite3.Connection, start: str, end: str, statuses: list[str], out_path: Path) -> int:
    rows = query_rows(conn, start, end, statuses)
    wb = Workbook()
    ws = wb.active
    ws.title = "入库明细"
    headers = [h for h, _ in COLUMNS]
    ws.append(headers)
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = _CENTER
    for r in rows:
        ws.append([fn(r) for _, fn in COLUMNS])
    for c, h in enumerate(headers, 1):
        ws.column_dimensions[get_column_letter(c)].width = max(10, min(24, len(h) * 2 + 8))
    ws.freeze_panes = "A2"
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return len(rows)


def default_export_name() -> str:
    return f"入库登记_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
